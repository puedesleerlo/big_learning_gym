"""Faithful parsing and anchored source versions precede interpretation."""

import csv
import io
import json
import re
import zipfile
from pathlib import Path

from .store import digest, now

PARSER_VERSION = "native-structure-v1"
MAX_BYTES = 25 * 1024 * 1024


def parse_file(name, raw):
    if len(raw) > MAX_BYTES:
        raise ValueError("Files are limited to 25 MB")
    suffix = Path(name).suffix.lower()
    fragments = []
    flags = []
    if suffix in {".docx", ".pptx"}:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded document exceeds the parsing limit")
    if suffix == ".pdf":
        from pypdf import PdfReader

        doc = PdfReader(io.BytesIO(raw))
        if len(doc.pages) > 500:
            raise ValueError("PDF is limited to 500 pages per upload")
        for i, page in enumerate(doc.pages, 1):
            text = page.extract_text() or ""
            fragments.append((f"page:{i}", text))
            if len(text.strip()) < 30:
                flags.append(f"page:{i}:needs_visual_review_or_ocr")
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(io.BytesIO(raw))
        for i, block in enumerate(doc.iter_inner_content(), 1):
            text = (
                block.text
                if isinstance(block, Paragraph)
                else "\n".join(" | ".join(c.text for c in row.cells) for row in block.rows)
                if isinstance(block, Table)
                else ""
            )
            fragments.append((f"block:{i}", text))
        flags.append("Comments, equations and drawing relationships require original-file review")
    elif suffix == ".pptx":
        from pptx import Presentation

        for i, slide in enumerate(Presentation(io.BytesIO(raw)).slides, 1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide:
                parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
            fragments.append((f"slide:{i}", "\n".join(parts)))
        flags.append("Diagrams require original-slide review")
    elif suffix == ".ipynb":
        for i, cell in enumerate(json.loads(raw).get("cells", []), 1):
            parts = [cell.get("cell_type", ""), "".join(cell.get("source", []))]
            for out in cell.get("outputs", []):
                parts.append("".join(out.get("text", out.get("data", {}).get("text/plain", []))))
            fragments.append((f"cell:{i}", "\n".join(parts)))
    elif suffix in {".csv", ".tsv"}:
        rows = csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter="\t" if suffix == ".tsv" else ",")
        fragments = [(f"row:{i}", " | ".join(row)) for i, row in enumerate(rows, 1)]
    elif suffix in {".md", ".txt", ".json", ".py", ".r", ".tex", ".html"}:
        text = raw.decode("utf-8-sig")
        if suffix == ".html":
            from html.parser import HTMLParser

            class TextParser(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts = []
                    self.hidden = 0

                def handle_starttag(self, tag, attrs):
                    if tag in {"script", "style"}:
                        self.hidden += 1

                def handle_endtag(self, tag):
                    if tag in {"script", "style"}:
                        self.hidden = max(0, self.hidden - 1)

                def handle_data(self, value):
                    if not self.hidden:
                        self.parts.append(value)

            parser = TextParser()
            parser.feed(text)
            text = "\n".join(parser.parts)
        lines = text.splitlines()
        start = 1
        chunk = []
        for i, line in enumerate(lines, 1):
            if chunk and (line.startswith("#") or sum(map(len, chunk)) > 4000):
                fragments.append((f"lines:{start}-{i - 1}", "\n".join(chunk)))
                chunk = []
                start = i
            chunk.append(line)
        if chunk:
            fragments.append((f"lines:{start}-{len(lines)}", "\n".join(chunk)))
    else:
        raise ValueError(
            "Supported: PDF, DOCX, PPTX, Markdown, text, HTML, JSON, CSV, TSV, code and notebooks"
        )
    # Keep long native sections anchored with an explicit part index.
    result = []
    for anchor, text in fragments:
        if not text.strip():
            continue
        for offset in range(0, len(text), 6000):
            result.append(
                {
                    "anchor": anchor + (f"/part:{offset // 6000 + 1}" if len(text) > 6000 else ""),
                    "text": text[offset : offset + 6000],
                }
            )
    if not result:
        raise ValueError("No usable text was extracted. Supply a text extraction or OCR this source first.")
    return result, flags


def ingest(store, course_id, name, raw, role="instruction"):
    if role not in {"instruction", "assessment", "rubric", "submission", "research"}:
        raise ValueError("Unknown source role")
    fragments, flags = parse_file(name, raw)
    name = Path(name).name
    checksum = digest(raw)
    source_id = "src_" + digest([course_id, name])[:24]
    version_id = "sv_" + digest([source_id, checksum])[:24]
    path = store.data_dir / "sources" / checksum
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_bytes(raw)
    with store.tx() as c:
        store.get(c, "course", course_id)
        previous = store.get(c, "source", version_id, required=False)
        if previous:
            return previous
        older = [s for s in store.list(c, "source") if s["source_id"] == source_id]
        result = store.put(
            c,
            "source",
            version_id,
            dict(
                source_id=source_id,
                course_id=course_id,
                name=name,
                content_hash=checksum,
                parser_version=PARSER_VERSION,
                created_at=now(),
                version=len(older) + 1,
                quality_flags=flags,
                fragment_count=len(fragments),
                storage_path=str(path.resolve()),
                latest=True,
                role=role,
                reconstruction_status="needs_review",
                previous_version_id=older[-1]["id"] if older else None,
            ),
        )
        for old in older:
            store.put(c, "source", old["id"], {**old, "latest": False})
        for i, fragment in enumerate(fragments):
            store.put(
                c,
                "fragment",
                f"{version_id}:{i}",
                {
                    **fragment,
                    "source_version_id": version_id,
                    "course_id": course_id,
                    "source_name": name,
                    "ordinal": i,
                },
            )
        store.emit(
            c,
            "source.ingested",
            version_id,
            {
                "course_id": course_id,
                "previous_version_ids": [s["id"] for s in older],
                "interpretation_status": "unconfirmed",
                "fragment_count": len(fragments),
            },
            key=f"source:{version_id}",
            source_id=source_id,
            source_revision=checksum,
            quality_flags=flags,
        )
        return result


def retrieve(store, c, course_id, query, source_ids=None, limit=12):
    terms = set(re.findall(r"\w{3,}", query.lower()))
    sources = {
        s["id"]
        for s in store.list(c, "source")
        if s["course_id"] == course_id and (s["id"] in source_ids if source_ids else s.get("latest", True))
    }
    candidates = [f for f in store.list(c, "fragment") if f["source_version_id"] in sources]

    def score(f):
        words = re.findall(r"\w+", f["text"].lower())
        return sum(min(words.count(t), 8) for t in terms) / (1 + len(words) / 1500)

    return sorted(candidates, key=score, reverse=True)[:limit]


def correct_reconstruction(store, source_id, fragments, expected_revision):
    if not fragments or any(not f.get("anchor") or not isinstance(f.get("text"), str) for f in fragments):
        raise ValueError("Corrections require anchored text fragments")
    if sum(len(f["text"]) for f in fragments) > 2_000_000:
        raise ValueError("Corrected text exceeds the document limit")
    with store.tx() as c:
        original = store.get(c, "source", source_id)
        if original["revision"] != expected_revision or not original["latest"]:
            raise ValueError("Source changed; review the latest extraction")
        ident = "sv_" + digest([source_id, fragments])[:24]
        previous = store.get(c, "source", ident, False)
        if previous:
            return previous
        store.put(c, "source", source_id, {**original, "latest": False})
        updated = store.put(
            c,
            "source",
            ident,
            {
                **original,
                "version": original["version"] + 1,
                "latest": True,
                "previous_version_id": source_id,
                "parser_version": "user-corrected-v1",
                "created_at": now(),
                "reconstruction_status": "confirmed",
                "fragment_count": len(fragments),
                "confirmed_at": now(),
                "quality_flags": ["Text corrected by learner; original file preserved"],
            },
        )
        for i, fragment in enumerate(fragments):
            store.put(
                c,
                "fragment",
                f"{ident}:{i}",
                {
                    "anchor": fragment["anchor"],
                    "text": fragment["text"],
                    "source_version_id": ident,
                    "course_id": original["course_id"],
                    "source_name": original["name"],
                    "ordinal": i,
                },
            )
        store.emit(
            c,
            "source.ingested",
            ident,
            {"course_id": original["course_id"], "previous_version_ids": [source_id], "correction": True},
        )
        return updated
